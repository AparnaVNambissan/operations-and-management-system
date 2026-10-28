import json
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.utils import timezone
from django.db.models import Sum, Q, Count
from django.contrib.auth import login as auth_login, logout as auth_logout

from .models import (
    Asset, AssetRequest, Category, Task, SubTask, Skill, Award, 
    User, Department, MaintenanceIssue, DepartmentIssue
)
from .forms import (
    CustomLoginForm, UserRegistrationForm, DepartmentForm, AssetForm, 
    CategoryForm, TaskForm, AwardForm, MaintenanceIssueForm, DepartmentIssueForm
)

# -------------------------------------------------------------
# Authentication & Access Control Views
# -------------------------------------------------------------

def custom_login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
        
    if request.method == 'POST':
        form = CustomLoginForm(request, data=request.POST)
        if form.is_valid():
            user = form.get_user()
            auth_login(request, user)
            messages.success(request, f'Welcome back, {user.full_name_or_username}!')
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid username or password. Please try again.')
    else:
        form = CustomLoginForm()
        
    return render(request, 'inventory/login.html', {'form': form})

def register(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
        
    if request.method == 'POST':
        form = UserRegistrationForm(request.POST)
        if form.is_valid():
            user = form.save()
            messages.success(request, 'Registration successful! You can now log in.')
            return redirect('login')
        else:
            messages.error(request, 'Please correct the errors below.')
    else:
        form = UserRegistrationForm()
        
    return render(request, 'inventory/register.html', {'form': form})

def custom_logout_view(request):
    auth_logout(request)
    messages.info(request, 'You have been logged out.')
    return redirect('login')

@login_required
def dashboard_view(request):
    if request.user.is_admin_role:
        return redirect('admin_dashboard')
    return redirect('employee_dashboard')

# -------------------------------------------------------------
# Admin Dashboard
# -------------------------------------------------------------

@login_required
def admin_dashboard(request):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')
    
    # KPIs
    total_assets_count = Asset.objects.count()
    low_stock_assets = [a for a in Asset.objects.select_related('category', 'department').all() if a.is_low_stock]
    pending_requests_count = AssetRequest.objects.filter(status='PENDING').count()
    active_tasks_count = Task.objects.exclude(status='COMPLETED').count()
    open_maintenance_count = MaintenanceIssue.objects.exclude(status='RESOLVED').count()
    open_dept_issues_count = DepartmentIssue.objects.exclude(status__in=['RESOLVED', 'CLOSED']).count()
    total_issues_count = open_maintenance_count + open_dept_issues_count
    total_employees_count = User.objects.filter(role='EMPLOYEE').count()
    total_depts_count = Department.objects.count()

    # Tables / lists
    pending_requests = AssetRequest.objects.filter(status='PENDING').select_related('employee', 'asset', 'employee__department').order_by('-request_date')[:10]
    recent_assets = Asset.objects.select_related('category', 'department').all().order_by('-id')[:8]
    departments = Department.objects.annotate(
        emp_count=Count('employees', distinct=True),
        asset_count=Count('assets', distinct=True),
        issue_count=Count('received_dept_issues', filter=~Q(received_dept_issues__status__in=['RESOLVED', 'CLOSED']), distinct=True)
    ).all()
    employees = User.objects.filter(role='EMPLOYEE').select_related('department').prefetch_related('tasks', 'awards').order_by('username')
    unassigned_tasks = Task.objects.filter(assigned_to__isnull=True).order_by('-created_at')[:5]
    maintenance_issues = MaintenanceIssue.objects.select_related('asset', 'department', 'reported_by').order_by('-reported_at')[:8]
    recent_awards = Award.objects.select_related('employee').order_by('-awarded_at')[:6]

    context = {
        'total_assets_count': total_assets_count,
        'low_stock_assets': low_stock_assets,
        'low_stock_count': len(low_stock_assets),
        'pending_requests_count': pending_requests_count,
        'active_tasks_count': active_tasks_count,
        'total_issues_count': total_issues_count,
        'total_employees_count': total_employees_count,
        'total_depts_count': total_depts_count,
        'pending_requests': pending_requests,
        'recent_assets': recent_assets,
        'departments': departments,
        'employees': employees,
        'unassigned_tasks': unassigned_tasks,
        'maintenance_issues': maintenance_issues,
        'recent_awards': recent_awards,
    }
    return render(request, 'inventory/admin_dashboard.html', context)

# -------------------------------------------------------------
# Employee Dashboard
# -------------------------------------------------------------

@login_required
def employee_dashboard(request):
    user = request.user
    department = user.department
    
    # Available assets
    all_assets = Asset.objects.select_related('category', 'department').all().order_by('name')
    
    # My requests
    my_requests = AssetRequest.objects.filter(employee=user).select_related('asset').order_by('-request_date')
    
    # My tasks
    my_tasks = Task.objects.filter(assigned_to=user).prefetch_related('subtasks').order_by('-created_at')
    
    # My awards
    my_awards = Award.objects.filter(employee=user).order_by('-awarded_at')
    
    # Issues
    all_departments = Department.objects.all()
    my_reported_maintenance = MaintenanceIssue.objects.filter(reported_by=user).select_related('asset', 'department').order_by('-reported_at')
    my_reported_dept_issues = DepartmentIssue.objects.filter(reported_by=user).select_related('target_department').order_by('-created_at')
    
    # Department incoming issues
    if department:
        received_maintenance_issues = MaintenanceIssue.objects.filter(department=department).select_related('asset', 'reported_by').order_by('-reported_at')
        received_dept_issues = DepartmentIssue.objects.filter(target_department=department).select_related('reported_by').order_by('-created_at')
    else:
        received_maintenance_issues = MaintenanceIssue.objects.none()
        received_dept_issues = DepartmentIssue.objects.none()
    
    # Stats
    active_tasks_count = my_tasks.exclude(status='COMPLETED').count()
    approved_requests_count = my_requests.filter(status='APPROVED').count()
    total_awards_count = my_awards.count()
    
    context = {
        'department': department,
        'all_assets': all_assets,
        'my_requests': my_requests,
        'my_tasks': my_tasks,
        'my_awards': my_awards,
        'all_departments': all_departments,
        'my_reported_maintenance': my_reported_maintenance,
        'my_reported_dept_issues': my_reported_dept_issues,
        'received_maintenance_issues': received_maintenance_issues,
        'received_dept_issues': received_dept_issues,
        'active_tasks_count': active_tasks_count,
        'approved_requests_count': approved_requests_count,
        'total_awards_count': total_awards_count,
    }
    return render(request, 'inventory/employee_dashboard.html', context)

# -------------------------------------------------------------
# Department Management (CRUD)
# -------------------------------------------------------------

@login_required
def department_list(request):
    search_query = request.GET.get('search', '').strip()
    departments = Department.objects.annotate(
        emp_count=Count('employees', distinct=True),
        asset_count=Count('assets', distinct=True),
        open_issues_count=Count('received_dept_issues', filter=~Q(received_dept_issues__status__in=['RESOLVED', 'CLOSED']), distinct=True),
        maintenance_count=Count('maintenance_issues', filter=~Q(maintenance_issues__status='RESOLVED'), distinct=True)
    )
    if search_query:
        departments = departments.filter(name__icontains=search_query)
        
    return render(request, 'inventory/department_list.html', {
        'departments': departments,
        'search_query': search_query
    })

@login_required
def add_department(request):
    if not request.user.is_admin_role:
        messages.error(request, 'Only administrators can add departments.')
        return redirect('department_list')
        
    if request.method == 'POST':
        form = DepartmentForm(request.POST)
        if form.is_valid():
            dept = form.save()
            messages.success(request, f'Department "{dept.name}" created successfully.')
            return redirect('department_list')
    else:
        form = DepartmentForm()
        
    return render(request, 'inventory/department_form.html', {'form': form, 'title': 'Add New Department'})

@login_required
def edit_department(request, dept_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Only administrators can edit departments.')
        return redirect('department_list')
        
    dept = get_object_or_404(Department, id=dept_id)
    if request.method == 'POST':
        form = DepartmentForm(request.POST, instance=dept)
        if form.is_valid():
            form.save()
            messages.success(request, f'Department "{dept.name}" updated successfully.')
            return redirect('department_list')
    else:
        form = DepartmentForm(instance=dept)
        
    return render(request, 'inventory/department_form.html', {'form': form, 'title': f'Edit Department: {dept.name}', 'dept': dept})

@login_required
def delete_department(request, dept_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Only administrators can delete departments.')
        return redirect('department_list')
        
    dept = get_object_or_404(Department, id=dept_id)
    if request.method == 'POST':
        name = dept.name
        # Reassign employees and assets safely to null
        dept.employees.update(department=None)
        dept.assets.update(department=None)
        dept.delete()
        messages.success(request, f'Department "{name}" has been deleted.')
        return redirect('department_list')
        
    return render(request, 'inventory/confirm_delete.html', {
        'item_type': 'Department',
        'item_name': dept.name,
        'cancel_url': 'department_list'
    })

@login_required
def department_detail(request, dept_id):
    dept = get_object_or_404(Department, id=dept_id)
    employees = dept.employees.all()
    assets = dept.assets.all()
    issues_received = dept.received_dept_issues.all().order_by('-created_at')
    maintenance_issues = dept.maintenance_issues.all().order_by('-reported_at')
    
    return render(request, 'inventory/department_detail.html', {
        'dept': dept,
        'employees': employees,
        'assets': assets,
        'issues_received': issues_received,
        'maintenance_issues': maintenance_issues
    })

# -------------------------------------------------------------
# Asset Management & Availability
# -------------------------------------------------------------

@login_required
def asset_list(request):
    search = request.GET.get('search', '').strip()
    category_id = request.GET.get('category')
    department_id = request.GET.get('department')
    stock_filter = request.GET.get('stock')

    assets = Asset.objects.select_related('category', 'department').all().order_by('name')

    if search:
        assets = assets.filter(Q(name__icontains=search) | Q(description__icontains=search))
    if category_id:
        assets = assets.filter(category_id=category_id)
    if department_id:
        assets = assets.filter(department_id=department_id)
    
    # Process stock filter
    asset_list_final = []
    for a in assets:
        if stock_filter == 'in_stock' and (a.available_quantity == 0 or a.is_low_stock):
            continue
        elif stock_filter == 'low_stock' and not (a.available_quantity > 0 and a.is_low_stock):
            continue
        elif stock_filter == 'out_of_stock' and a.available_quantity > 0:
            continue
        asset_list_final.append(a)

    categories = Category.objects.all()
    departments = Department.objects.all()

    return render(request, 'inventory/asset_list.html', {
        'assets': asset_list_final,
        'categories': categories,
        'departments': departments,
        'search': search,
        'selected_category': category_id,
        'selected_department': department_id,
        'selected_stock': stock_filter,
    })

@login_required
def add_asset(request):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('employee_dashboard')
        
    if request.method == 'POST':
        form = AssetForm(request.POST)
        if form.is_valid():
            asset = form.save()
            messages.success(request, f'Asset "{asset.name}" added successfully.')
            return redirect('asset_list')
    else:
        form = AssetForm()
    
    return render(request, 'inventory/asset_form.html', {'form': form, 'title': 'Add New Asset'})

@login_required
def edit_asset(request, asset_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('employee_dashboard')
        
    asset = get_object_or_404(Asset, id=asset_id)
    if request.method == 'POST':
        form = AssetForm(request.POST, instance=asset)
        if form.is_valid():
            form.save()
            messages.success(request, f'Asset "{asset.name}" updated successfully.')
            return redirect('asset_list')
    else:
        form = AssetForm(instance=asset)
    
    return render(request, 'inventory/asset_form.html', {'form': form, 'title': f'Edit Asset: {asset.name}', 'asset': asset})

@login_required
def delete_asset(request, asset_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('asset_list')
        
    asset = get_object_or_404(Asset, id=asset_id)
    if request.method == 'POST':
        name = asset.name
        asset.delete()
        messages.success(request, f'Asset "{name}" has been removed.')
        return redirect('asset_list')
        
    return render(request, 'inventory/confirm_delete.html', {
        'item_type': 'Asset',
        'item_name': asset.name,
        'cancel_url': 'asset_list'
    })

@login_required
def category_list(request):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('asset_list')
        
    if request.method == 'POST':
        form = CategoryForm(request.POST)
        if form.is_valid():
            cat = form.save()
            messages.success(request, f'Category "{cat.name}" added.')
            return redirect('category_list')
    else:
        form = CategoryForm()
        
    categories = Category.objects.annotate(asset_count=Count('assets')).all()
    return render(request, 'inventory/category_list.html', {'categories': categories, 'form': form})

@login_required
def delete_category(request, cat_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('category_list')
        
    category = get_object_or_404(Category, id=cat_id)
    if request.method == 'POST':
        name = category.name
        category.delete()
        messages.success(request, f'Category "{name}" deleted.')
        return redirect('category_list')
        
    return render(request, 'inventory/confirm_delete.html', {
        'item_type': 'Category',
        'item_name': category.name,
        'cancel_url': 'category_list'
    })

@login_required
def add_category_api(request):
    if not request.user.is_admin_role:
        return JsonResponse({'error': 'Unauthorized'}, status=403)
        
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            name = data.get('name', '').strip()
            description = data.get('description', '').strip()
            if name:
                category, created = Category.objects.get_or_create(name=name, defaults={'description': description})
                return JsonResponse({'id': category.id, 'name': category.name})
            return JsonResponse({'error': 'Name is required'}, status=400)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=400)
    return JsonResponse({'error': 'Invalid method'}, status=405)

# -------------------------------------------------------------
# Asset Requests, Approvals, Rejections, and Returns
# -------------------------------------------------------------

@login_required
def request_list(request):
    status_filter = request.GET.get('status', '').upper()
    search = request.GET.get('search', '').strip()
    
    if request.user.is_admin_role:
        requests_qs = AssetRequest.objects.select_related('employee', 'asset', 'employee__department').all().order_by('-request_date')
    else:
        requests_qs = AssetRequest.objects.filter(employee=request.user).select_related('asset').order_by('-request_date')
        
    if status_filter:
        requests_qs = requests_qs.filter(status=status_filter)
    if search:
        requests_qs = requests_qs.filter(Q(asset__name__icontains=search) | Q(employee__username__icontains=search))
        
    return render(request, 'inventory/request_list.html', {
        'requests': requests_qs,
        'selected_status': status_filter,
        'search': search
    })

@login_required
def request_asset(request, asset_id):
    asset = get_object_or_404(Asset, id=asset_id)
    if request.method == 'POST':
        try:
            quantity = int(request.POST.get('quantity', 1))
        except ValueError:
            quantity = 1
        notes = request.POST.get('notes', '').strip()
        
        if quantity <= 0:
            messages.error(request, 'Quantity must be at least 1.')
        elif quantity > asset.available_quantity:
            messages.error(request, f'Cannot request {quantity}. Only {asset.available_quantity} available.')
        else:
            AssetRequest.objects.create(
                employee=request.user,
                asset=asset,
                quantity=quantity,
                notes=notes,
                status='PENDING'
            )
            messages.success(request, f'Request for {quantity}x "{asset.name}" submitted successfully.')
            
    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or 'employee_dashboard'
    return redirect(next_url)

@login_required
def approve_request(request, request_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('employee_dashboard')
        
    asset_req = get_object_or_404(AssetRequest, id=request_id)
    if asset_req.status == 'PENDING':
        if asset_req.asset.available_quantity >= asset_req.quantity:
            asset_req.asset.available_quantity -= asset_req.quantity
            asset_req.asset.save()
            asset_req.status = 'APPROVED'
            asset_req.response_date = timezone.now()
            asset_req.save()
            messages.success(request, f'Request #{asset_req.id} for {asset_req.employee.username} approved.')
        else:
            messages.error(request, f'Insufficient stock! Available: {asset_req.asset.available_quantity}, Requested: {asset_req.quantity}.')
    else:
        messages.warning(request, 'Request is not in PENDING state.')
        
    next_url = request.META.get('HTTP_REFERER') or 'admin_dashboard'
    return redirect(next_url)

@login_required
def reject_request(request, request_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('employee_dashboard')
        
    asset_req = get_object_or_404(AssetRequest, id=request_id)
    if asset_req.status == 'PENDING':
        reason = request.POST.get('admin_notes', 'Rejected by administrator.').strip()
        asset_req.status = 'REJECTED'
        asset_req.admin_notes = reason
        asset_req.response_date = timezone.now()
        asset_req.save()
        messages.info(request, f'Request #{asset_req.id} rejected.')
        
    next_url = request.META.get('HTTP_REFERER') or 'admin_dashboard'
    return redirect(next_url)

@login_required
def return_asset(request, request_id):
    asset_req = get_object_or_404(AssetRequest, id=request_id)
    
    # Only the borrower or an admin can mark return
    if not request.user.is_admin_role and request.user != asset_req.employee:
        messages.error(request, 'Unauthorized.')
        return redirect('employee_dashboard')
        
    if asset_req.status == 'APPROVED':
        asset_req.asset.available_quantity += asset_req.quantity
        if asset_req.asset.available_quantity > asset_req.asset.total_quantity:
            asset_req.asset.available_quantity = asset_req.asset.total_quantity
        asset_req.asset.save()
        
        asset_req.status = 'RETURNED'
        asset_req.return_date = timezone.now()
        asset_req.save()
        messages.success(request, f'Asset "{asset_req.asset.name}" returned. Inventory stock restored.')
    else:
        messages.warning(request, 'Only approved assets can be marked as returned.')
        
    next_url = request.META.get('HTTP_REFERER') or 'dashboard'
    return redirect(next_url)

# -------------------------------------------------------------
# Complaints & Issues Management
# -------------------------------------------------------------

@login_required
def issue_list(request):
    tab = request.GET.get('tab', 'maintenance')
    status_filter = request.GET.get('status', '').upper()
    priority_filter = request.GET.get('priority', '').upper()
    dept_id = request.GET.get('dept')

    maintenance_issues = MaintenanceIssue.objects.select_related('asset', 'department', 'reported_by').all().order_by('-reported_at')
    dept_issues = DepartmentIssue.objects.select_related('target_department', 'reported_by', 'reported_by__department').all().order_by('-created_at')

    # Apply filters
    if status_filter:
        maintenance_issues = maintenance_issues.filter(status=status_filter)
        dept_issues = dept_issues.filter(status=status_filter)
    if priority_filter:
        maintenance_issues = maintenance_issues.filter(priority=priority_filter)
        dept_issues = dept_issues.filter(priority=priority_filter)
    if dept_id:
        maintenance_issues = maintenance_issues.filter(department_id=dept_id)
        dept_issues = dept_issues.filter(target_department_id=dept_id)

    departments = Department.objects.all()
    assets = Asset.objects.all()

    return render(request, 'inventory/issue_list.html', {
        'tab': tab,
        'maintenance_issues': maintenance_issues,
        'dept_issues': dept_issues,
        'departments': departments,
        'assets': assets,
        'status_filter': status_filter,
        'priority_filter': priority_filter,
        'selected_dept': dept_id,
    })

@login_required
def report_issue(request):
    if request.method == 'POST':
        asset_id = request.POST.get('asset_id')
        description = request.POST.get('description', '').strip()
        is_damaged = request.POST.get('is_damaged') == 'on'
        priority = request.POST.get('priority', 'MEDIUM')
        
        asset = get_object_or_404(Asset, id=asset_id)
        target_department = asset.department
        
        MaintenanceIssue.objects.create(
            asset=asset,
            department=target_department,
            reported_by=request.user,
            description=description,
            is_damaged=is_damaged,
            priority=priority,
            status='REPORTED'
        )
        messages.success(request, f'Issue on "{asset.name}" reported successfully.')
        
    next_url = request.META.get('HTTP_REFERER') or 'employee_dashboard'
    return redirect(next_url)

@login_required
def create_dept_issue(request):
    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        description = request.POST.get('description', '').strip()
        target_dept_id = request.POST.get('target_department')
        priority = request.POST.get('priority', 'MEDIUM')
        
        target_dept = get_object_or_404(Department, id=target_dept_id)
        
        DepartmentIssue.objects.create(
            title=title,
            description=description,
            target_department=target_dept,
            reported_by=request.user,
            priority=priority,
            status='OPEN'
        )
        messages.success(request, f'Department issue "{title}" submitted to {target_dept.name}.')
        
    next_url = request.META.get('HTTP_REFERER') or 'employee_dashboard'
    return redirect(next_url)

@login_required
def update_maintenance_issue(request, issue_id):
    issue = get_object_or_404(MaintenanceIssue, id=issue_id)
    
    # Permission: Admin or members of the responsible department
    can_update = request.user.is_admin_role or (issue.department and request.user.department == issue.department)
    if not can_update:
        messages.error(request, 'You do not have permission to update this issue.')
        return redirect('issue_list')
        
    if request.method == 'POST':
        status = request.POST.get('status')
        priority = request.POST.get('priority')
        resolution_notes = request.POST.get('resolution_notes', '').strip()
        
        if status in dict(MaintenanceIssue.STATUS_CHOICES):
            issue.status = status
            if status == 'RESOLVED':
                issue.resolved_at = timezone.now()
            else:
                issue.resolved_at = None
        if priority in dict(MaintenanceIssue.PRIORITY_CHOICES):
            issue.priority = priority
        if resolution_notes:
            issue.resolution_notes = resolution_notes
            
        issue.save()
        messages.success(request, f'Maintenance issue #{issue.id} status updated to {issue.status}.')
        
    next_url = request.META.get('HTTP_REFERER') or 'issue_list'
    return redirect(next_url)

@login_required
def update_dept_issue(request, issue_id):
    issue = get_object_or_404(DepartmentIssue, id=issue_id)
    
    # Permission: Admin or members of target department
    can_update = request.user.is_admin_role or request.user.department == issue.target_department
    if not can_update:
        messages.error(request, 'You do not have permission to update this issue.')
        return redirect('issue_list')
        
    if request.method == 'POST':
        status = request.POST.get('status')
        priority = request.POST.get('priority')
        resolution_notes = request.POST.get('resolution_notes', '').strip()
        
        if status in dict(DepartmentIssue.STATUS_CHOICES):
            issue.status = status
            if status in ['RESOLVED', 'CLOSED']:
                issue.resolved_at = timezone.now()
            else:
                issue.resolved_at = None
        if priority in dict(DepartmentIssue.PRIORITY_CHOICES):
            issue.priority = priority
        if resolution_notes:
            issue.resolution_notes = resolution_notes
            
        issue.save()
        messages.success(request, f'Department issue #{issue.id} updated to {issue.status}.')
        
    next_url = request.META.get('HTTP_REFERER') or 'issue_list'
    return redirect(next_url)

@login_required
def delete_issue(request, issue_type, issue_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('issue_list')
        
    if issue_type == 'maintenance':
        issue = get_object_or_404(MaintenanceIssue, id=issue_id)
    else:
        issue = get_object_or_404(DepartmentIssue, id=issue_id)
        
    if request.method == 'POST':
        issue.delete()
        messages.success(request, 'Issue deleted successfully.')
        return redirect('issue_list')
        
    return render(request, 'inventory/confirm_delete.html', {
        'item_type': 'Issue',
        'item_name': str(issue),
        'cancel_url': 'issue_list'
    })

# -------------------------------------------------------------
# Task & Subtask Management
# -------------------------------------------------------------

@login_required
def task_list(request):
    status_filter = request.GET.get('status', '').upper()
    priority_filter = request.GET.get('priority', '').upper()
    employee_id = request.GET.get('employee')

    if request.user.is_admin_role:
        tasks = Task.objects.select_related('assigned_to', 'required_skill').prefetch_related('subtasks').all().order_by('-created_at')
    else:
        tasks = Task.objects.filter(assigned_to=request.user).select_related('required_skill').prefetch_related('subtasks').order_by('-created_at')

    if status_filter:
        tasks = tasks.filter(status=status_filter)
    if priority_filter:
        tasks = tasks.filter(priority=priority_filter)
    if employee_id and request.user.is_admin_role:
        tasks = tasks.filter(assigned_to_id=employee_id)

    employees = User.objects.filter(role='EMPLOYEE')

    return render(request, 'inventory/task_list.html', {
        'tasks': tasks,
        'employees': employees,
        'status_filter': status_filter,
        'priority_filter': priority_filter,
        'selected_employee': employee_id,
    })

@login_required
def create_task(request):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')

    if request.method == 'POST':
        form = TaskForm(request.POST)
        if form.is_valid():
            task = form.save(commit=False)
            
            # Auto-assign algorithm
            if 'auto_assign' in request.POST and task.required_skill:
                eligible_employees = User.objects.filter(role='EMPLOYEE', skills=task.required_skill)
                if eligible_employees.exists():
                    eligible_employees = eligible_employees.annotate(
                        current_workload=Sum('tasks__workload_hours', filter=Q(tasks__status__in=['PENDING', 'IN_PROGRESS']))
                    )
                    best_employee = None
                    lowest_workload = float('inf')
                    for emp in eligible_employees:
                        wl = emp.current_workload or 0
                        if wl < lowest_workload:
                            lowest_workload = wl
                            best_employee = emp
                    if best_employee:
                        task.assigned_to = best_employee
                        messages.success(request, f'Task auto-assigned to {best_employee.username} (Current workload: {lowest_workload}h).')
                    else:
                        messages.warning(request, 'Could not determine best employee for auto-assignment.')
                else:
                    messages.warning(request, 'No employee with the required skill found. Task created unassigned.')
            
            task.save()
            
            # Subtasks
            subtask_titles = request.POST.getlist('subtasks')
            for title in subtask_titles:
                if title.strip():
                    SubTask.objects.create(task=task, title=title.strip())
                    
            messages.success(request, f'Task "{task.title}" created successfully.')
            return redirect('task_list')
    else:
        initial = {}
        employee_id = request.GET.get('employee')
        if employee_id:
            initial['assigned_to'] = employee_id
        form = TaskForm(initial=initial)
    
    return render(request, 'inventory/task_form.html', {'form': form, 'title': 'Create New Task'})

@login_required
def edit_task(request, task_id):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')
        
    task = get_object_or_404(Task, id=task_id)
    if request.method == 'POST':
        form = TaskForm(request.POST, instance=task)
        if form.is_valid():
            form.save()
            # Add any new subtasks submitted
            new_subtasks = request.POST.getlist('new_subtasks')
            for st in new_subtasks:
                if st.strip():
                    SubTask.objects.create(task=task, title=st.strip())
                    
            messages.success(request, f'Task "{task.title}" updated.')
            return redirect('task_list')
    else:
        form = TaskForm(instance=task)
        
    return render(request, 'inventory/task_form.html', {'form': form, 'title': f'Edit Task: {task.title}', 'task': task})

@login_required
def delete_task(request, task_id):
    if not request.user.is_admin_role:
        return redirect('task_list')
        
    task = get_object_or_404(Task, id=task_id)
    if request.method == 'POST':
        name = task.title
        task.delete()
        messages.success(request, f'Task "{name}" deleted.')
        return redirect('task_list')
        
    return render(request, 'inventory/confirm_delete.html', {
        'item_type': 'Task',
        'item_name': task.title,
        'cancel_url': 'task_list'
    })

@login_required
def update_subtask_status(request, subtask_id):
    subtask = get_object_or_404(SubTask, id=subtask_id)
    task = subtask.task
    
    if request.user != task.assigned_to and not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('employee_dashboard')
        
    if request.method == 'POST':
        is_completed = request.POST.get('is_completed') == 'on'
        subtask.is_completed = is_completed
        subtask.save()
        
        subtasks_list = list(task.subtasks.all())
        all_completed = len(subtasks_list) > 0 and all(st.is_completed for st in subtasks_list)
        any_completed = any(st.is_completed for st in subtasks_list)
        
        if all_completed and task.status != 'COMPLETED':
            task.status = 'COMPLETED'
            task.completed_at = timezone.now()
            task.save()
            
            # Gamification Award for timely completion
            if task.completed_at <= task.deadline and task.assigned_to:
                Award.objects.create(
                    employee=task.assigned_to,
                    badge_icon='🏆',
                    name=f"Timely Finish: {task.title[:30]}",
                    description=f"Completed '{task.title}' on time before deadline."
                )
                messages.success(request, f"🎉 Congratulations! You completed '{task.title}' on time and earned an achievement badge!")
            else:
                messages.success(request, f"Task '{task.title}' marked as completed.")
        elif any_completed and task.status != 'IN_PROGRESS':
            task.status = 'IN_PROGRESS'
            task.completed_at = None
            task.save()
        elif not any_completed and task.status != 'PENDING':
            task.status = 'PENDING'
            task.completed_at = None
            task.save()
            
    next_url = request.META.get('HTTP_REFERER') or ('employee_dashboard' if not request.user.is_admin_role else 'admin_dashboard')
    return redirect(next_url)

# -------------------------------------------------------------
# Awards & Badge Management
# -------------------------------------------------------------

@login_required
def award_list(request):
    awards = Award.objects.select_related('employee', 'employee__department').all().order_by('-awarded_at')
    
    # Leaderboard / Top badge earners
    leaderboard = User.objects.filter(role='EMPLOYEE').annotate(
        award_count=Count('awards')
    ).filter(award_count__gt=0).order_by('-award_count')[:10]
    
    return render(request, 'inventory/award_list.html', {
        'awards': awards,
        'leaderboard': leaderboard,
        'badge_choices': Award.BADGE_CHOICES
    })

@login_required
def grant_award(request):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')

    if request.method == 'POST':
        form = AwardForm(request.POST)
        if form.is_valid():
            award = form.save()
            messages.success(request, f'Badge "{award.name}" granted to {award.employee.username}!')
            return redirect('award_list')
    else:
        initial = {}
        employee_id = request.GET.get('employee')
        if employee_id:
            initial['employee'] = employee_id
        form = AwardForm(initial=initial)
        
    return render(request, 'inventory/award_form.html', {'form': form, 'title': 'Grant Custom Award / Badge'})

@login_required
def delete_award(request, award_id):
    if not request.user.is_admin_role:
        messages.error(request, 'Unauthorized.')
        return redirect('award_list')
        
    award = get_object_or_404(Award, id=award_id)
    if request.method == 'POST':
        name = award.name
        award.delete()
        messages.success(request, f'Award "{name}" deleted.')
        return redirect('award_list')
        
    return render(request, 'inventory/confirm_delete.html', {
        'item_type': 'Award',
        'item_name': f"{award.badge_icon} {award.name} ({award.employee.username})",
        'cancel_url': 'award_list'
    })

# -------------------------------------------------------------
# Employee / Team Management (Admin)
# -------------------------------------------------------------

@login_required
def employee_list(request):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')
        
    employees = User.objects.select_related('department').prefetch_related('skills', 'tasks', 'awards').all().order_by('username')
    departments = Department.objects.all()
    
    return render(request, 'inventory/employee_list.html', {
        'employees': employees,
        'departments': departments
    })

@login_required
def update_user_department(request, user_id):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')
        
    target_user = get_object_or_404(User, id=user_id)
    if request.method == 'POST':
        dept_id = request.POST.get('department')
        role = request.POST.get('role')
        if dept_id:
            target_user.department = get_object_or_404(Department, id=dept_id)
        else:
            target_user.department = None
        if role in ['ADMIN', 'EMPLOYEE']:
            target_user.role = role
        target_user.save()
        messages.success(request, f'Updated {target_user.username}\'s profile.')
        
    return redirect('employee_list')

# -------------------------------------------------------------
# Reports
# -------------------------------------------------------------

@login_required
def generate_pdf_report(request):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')
    import io
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from django.http import FileResponse
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('ReportTitle', parent=styles['Heading1'], fontSize=20, textColor=colors.HexColor('#1E3A8A'))
    elements.append(Paragraph("SecureCorp Inventory & Asset Report", title_style))
    elements.append(Paragraph(f"Generated on: {timezone.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 16))
    
    # Table data
    data = [["Asset Name", "Category", "Department", "Total", "Available", "Status"]]
    for asset in Asset.objects.select_related('category', 'department').all().order_by('name'):
        status = "Low Stock" if asset.is_low_stock else ("Out of Stock" if asset.available_quantity == 0 else "In Stock")
        data.append([
            asset.name,
            asset.category.name if asset.category else "—",
            asset.department.name if asset.department else "General",
            str(asset.total_quantity),
            str(asset.available_quantity),
            status
        ])
        
    table = Table(data, colWidths=[150, 80, 100, 50, 60, 90])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
    ]))
    elements.append(table)
    
    doc.build(elements)
    buffer.seek(0)
    return FileResponse(buffer, as_attachment=True, filename=f'inventory_report_{timezone.now().strftime("%Y%m%d")}.pdf')

@login_required
def generate_excel_report(request):
    if not request.user.is_admin_role:
        return redirect('employee_dashboard')
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=inventory_report_{timezone.now().strftime("%Y%m%d")}.xlsx'
    
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = 'Inventory'
    
    headers = ['ID', 'Asset Name', 'Category', 'Department', 'Total Quantity', 'Available Quantity', 'Low Stock Threshold', 'Status']
    worksheet.append(headers)
    
    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        
    for asset in Asset.objects.select_related('category', 'department').all().order_by('id'):
        status = "Out of Stock" if asset.available_quantity == 0 else ("Low Stock" if asset.is_low_stock else "In Stock")
        worksheet.append([
            asset.id,
            asset.name,
            asset.category.name if asset.category else '',
            asset.department.name if asset.department else 'General',
            asset.total_quantity,
            asset.available_quantity,
            asset.low_stock_threshold,
            status
        ])
        
    for col in worksheet.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    workbook.save(response)
    return response
